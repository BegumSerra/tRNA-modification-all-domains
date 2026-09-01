import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill
import sys


def color_presence_values(data_file, color_file, output_file):
    # Load the data file
    df = pd.read_csv(data_file, sep='\t')
    print(f"Data shape: {df.shape}")
    print(f"Columns: {df.columns.tolist()}")
    print(f"First few rows:\n{df.head()}")

    # Load the color order from the color file (no header, one color per row)
    with open(color_file, 'r') as f:
        colors = [line.strip() for line in f.readlines()]
    
    print(f"Colors loaded: {len(colors)}")
    print(f"First few colors: {colors[:5]}")

    data_rows = len(df)   
    if len(colors) != data_rows:
        raise ValueError(
            f"The number of colors ({len(colors)}) does not match the number of data rows ({data_rows}) in the file.")

    # Create a new Excel workbook and worksheet
    wb = Workbook()
    ws = wb.active
    ws.title = "Colored Table"

    # Write the header row
    for col_idx, column_name in enumerate(df.columns, start=1):
        header_cell = ws.cell(row=1, column=col_idx, value=column_name)
        header_cell.fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    # Identify protein columns (exclude taxonomy columns)
    taxonomy_cols = ['Domain', 'Phylum', 'Order', 'Family', 'Species']
    protein_cols = [col for col in df.columns if col not in taxonomy_cols]
    print(f"Protein columns: {protein_cols}")

    # Write the data rows and color presence values
    colors_applied = 0
    for row_idx, (index, row) in enumerate(df.iterrows(), start=2):  # Start at 2 to skip header
        color = colors[row_idx - 2] if row_idx - 2 < len(colors) else None  # Get the color for the row
        
        print(f"Processing row {row_idx}, color: {color}")
        
        for col_idx, (col_name, value) in enumerate(row.items(), start=1):
            cell = ws.cell(row=row_idx, column=col_idx)

            # Set the cell value
            if pd.isna(value) or value == 'N/A':
                cell.value = 'N/A'
            else:
                cell.value = value

            # Apply color only to protein columns with presence values
            if (col_name in protein_cols and 
                color and 
                str(row.get('Species', '')).strip() != 'Total'):
                
                # Check if value is 0 or 1
                try:
                    numeric_value = float(value) if not pd.isna(value) else None
                    if numeric_value in [0.0, 1.0]:
                        # Clean and validate color
                        clean_color = color.replace('#', '').strip().upper()
                        
                        if len(clean_color) == 6 and all(c in '0123456789ABCDEF' for c in clean_color):
                            fill = PatternFill(start_color=clean_color, end_color=clean_color, fill_type="solid")
                            cell.fill = fill
                            colors_applied += 1
                            print(f"Applied color {clean_color} to cell ({row_idx},{col_idx}) - {col_name}={value}")
                        else:
                            print(f"Invalid color format: {clean_color}")
                except (ValueError, TypeError):
                    print(f"Could not convert value to float: {value}")

    print(f"Total colors applied: {colors_applied}")
    
    # Save the workbook
    wb.save(output_file)
    print(f"Colored table saved to '{output_file}'.")


if __name__ == "__main__":
    data_file = sys.argv[1]  # Input data file
    color_file = sys.argv[2]  # Text file with color order
    output_file = sys.argv[3]  # Output Excel file

    color_presence_values(data_file, color_file, output_file)