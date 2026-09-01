#!/usr/bin/env python3
"""
Order and color presence/absence table based on taxa_color_manual_order.tsv.
Groups species by the 'Group' column and colors rows based on 'Colorcode'.
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Border, Side, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows
import sys


def order_and_color_table(data_file, order_file, output_file):
    """
    Order the presence/absence table based on taxa_color_manual_order.tsv
    and color rows based on Group/Colorcode.
    
    Args:
        data_file: Path to presence/absence TSV file
        order_file: Path to taxa_color_manual_order.tsv (with Group and Colorcode columns)
        output_file: Path to output Excel file
    """
    # Load the data file
    df = pd.read_csv(data_file, sep='\t')
    print(f"Data shape: {df.shape}")
    print(f"Data columns: {df.columns.tolist()[:10]}...")
    
    # Load the order file with Group and Colorcode
    order_df = pd.read_csv(order_file, sep='\t')
    print(f"Order file shape: {order_df.shape}")
    print(f"Order file columns: {order_df.columns.tolist()}")
    
    # Create a mapping from Species to (Group, Colorcode, order_index)
    species_info = {}
    for idx, row in order_df.iterrows():
        species = row['Species']
        group = row['Group']
        colorcode = row['Colorcode']
        species_info[species] = {
            'Group': group,
            'Colorcode': colorcode,
            'order_index': idx
        }
    
    print(f"Species in order file: {len(species_info)}")
    
    # Add Group, Colorcode, and order_index to data
    df['Group'] = df['Species'].map(lambda x: species_info.get(x, {}).get('Group', 'Unknown'))
    df['Colorcode'] = df['Species'].map(lambda x: species_info.get(x, {}).get('Colorcode', '#FFFFFF'))
    df['order_index'] = df['Species'].map(lambda x: species_info.get(x, {}).get('order_index', 99999))
    
    # Check for missing species
    missing = df[df['order_index'] == 99999]['Species'].unique()
    if len(missing) > 0:
        print(f"Warning: {len(missing)} species not found in order file:")
        for s in missing[:10]:
            print(f"  - {s}")
        if len(missing) > 10:
            print(f"  ... and {len(missing) - 10} more")
    
    # Sort by order_index (which preserves the order from order file)
    df = df.sort_values('order_index').reset_index(drop=True)
    
    # Keep Group column and insert it after Domain
    group_col = df['Group'].copy()
    colorcode_col = df['Colorcode'].copy()
    df = df.drop(columns=['Colorcode', 'order_index'])

    # Reorder columns to place Group right after Domain
    cols = list(df.columns)
    cols.remove('Group')
    domain_idx = cols.index('Domain')
    cols.insert(domain_idx + 1, 'Group')
    df = df[cols]
    
    # Create Excel workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Presence_Absence"
    
    # Define styles
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Identify taxonomy and protein columns
    taxonomy_cols = ['Domain', 'Group', 'Phylum', 'Order', 'Family', 'Species']
    protein_cols = [col for col in df.columns if col not in taxonomy_cols]
    
    # Write header row
    for col_idx, column_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=column_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center')
    
    # Track current group for adding separators
    current_group = None
    row_offset = 2  # Start at row 2 (after header)
    
    # Write data rows with coloring
    for idx, (_, row) in enumerate(df.iterrows()):
        group = group_col.iloc[idx]
        colorcode = colorcode_col.iloc[idx]
        
        # Add group separator row if group changes
        if group != current_group and current_group is not None:
            # Add an empty separator row
            for col_idx in range(1, len(df.columns) + 1):
                sep_cell = ws.cell(row=row_offset, column=col_idx, value="")
                sep_cell.fill = PatternFill(start_color="000000", end_color="000000", fill_type="solid")
            row_offset += 1
        
        current_group = group
        
        # Clean color code
        clean_color = colorcode.replace('#', '').strip().upper() if isinstance(colorcode, str) else 'FFFFFF'
        if len(clean_color) != 6 or not all(c in '0123456789ABCDEF' for c in clean_color):
            clean_color = 'FFFFFF'
        
        # Write the row
        for col_idx, (col_name, value) in enumerate(row.items(), start=1):
            cell = ws.cell(row=row_offset, column=col_idx)
            
            # Set cell value
            if pd.isna(value) or value == 'N/A':
                cell.value = 'N/A'
            else:
                cell.value = value
            
            cell.border = thin_border
            
            # Apply color to protein columns
            if col_name in protein_cols:
                try:
                    # Check if it's a presence value (0, 1, or numeric)
                    numeric_value = float(value) if not pd.isna(value) and str(value).lower() not in ['false', 'true', 'n/a'] else None
                    if numeric_value is not None or str(value).lower() in ['false', 'true']:
                        fill = PatternFill(start_color=clean_color, end_color=clean_color, fill_type="solid")
                        cell.fill = fill
                except (ValueError, TypeError):
                    # If value is FALSE/TRUE string
                    if str(value).upper() in ['FALSE', 'TRUE']:
                        fill = PatternFill(start_color=clean_color, end_color=clean_color, fill_type="solid")
                        cell.fill = fill
        
        row_offset += 1
    
    # Adjust column widths
    for col_idx, column_name in enumerate(df.columns, start=1):
        if column_name == 'Species':
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = 40
        elif column_name in taxonomy_cols:
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = 20
        else:
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = 12
    
    # Freeze the header row and taxonomy columns (Domain, Group, Phylum, Order, Family, Species = 6 cols -> G2)
    ws.freeze_panes = 'G2'
    
    # Save the workbook
    wb.save(output_file)
    print(f"Output saved to: {output_file}")
    print(f"Total rows: {row_offset - 2}")


if __name__ == "__main__":
    if len(sys.argv) != 4:
        print("Usage: python order-color-PA-by-group.py <data_file> <order_file> <output_file>")
        print("  data_file: Presence/absence TSV file")
        print("  order_file: taxa_color_manual_order.tsv with Group and Colorcode columns")
        print("  output_file: Output Excel file")
        sys.exit(1)
    
    data_file = sys.argv[1]
    order_file = sys.argv[2]
    output_file = sys.argv[3]
    
    order_and_color_table(data_file, order_file, output_file)
