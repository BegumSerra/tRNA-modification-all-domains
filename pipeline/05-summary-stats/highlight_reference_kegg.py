#!/usr/bin/env python3
"""
Highlight rows in protein tables that have the KEGG from protein_kegg.tsv.
For each protein table:
1. Look up the correct KEGG ID from protein_kegg.tsv
2. Highlight all rows with matching KEGG in blue

Author: GitHub Copilot
Date: 2025-11-28
"""

import pandas as pd
import argparse
import os
import re
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from pathlib import Path


def load_protein_kegg_mapping(protein_kegg_file):
    """
    Load protein to KEGG mapping from protein_kegg.tsv.
    Returns dict: {protein_name: kegg_id}
    """
    protein_kegg = {}
    with open(protein_kegg_file, 'r') as f:
        for line in f:
            line = line.strip()
            if line and ',' in line:
                parts = line.split(',')
                if len(parts) >= 2:
                    protein_name = parts[0].strip()
                    kegg_id = parts[1].strip()
                    protein_kegg[protein_name] = kegg_id
    return protein_kegg


def parse_fasta_headers(ref_seq_dirs):
    """
    Parse reference sequence fasta files to extract species names.
    Returns dict: {protein_name: {species: protein_id}}
    """
    ref_sequences = {}
    
    for domain_dir in ref_seq_dirs:
        if not os.path.exists(domain_dir):
            continue
            
        for fasta_file in Path(domain_dir).glob('*.fa'):
            protein_name = fasta_file.stem  # e.g., ADAT2 from ADAT2.fa
            
            if protein_name not in ref_sequences:
                ref_sequences[protein_name] = {}
            
            with open(fasta_file, 'r') as f:
                for line in f:
                    if line.startswith('>'):
                        # Extract protein ID and species from header
                        # Format examples:
                        # >WP_012899688.1 nucleoside deaminase [Methanocella paludicola]
                        # >sp|Q7Z6V5|ADAT2_HUMAN tRNA-specific adenosine deaminase 2 OS=Homo sapiens ...
                        header = line.strip()[1:]
                        
                        # Try UniProt format: OS=Species name
                        os_match = re.search(r'OS=([^=]+?)(?:\s+OX=|\s+GN=|$)', header)
                        if os_match:
                            species_full = os_match.group(1).strip()
                            # Extract only genus and species (first two words before parenthesis or strain)
                            species_match = re.match(r'(\S+\s+\S+)', species_full)
                            if species_match:
                                species = species_match.group(1).strip()
                            else:
                                species = species_full
                            protein_id = header.split()[0].split('|')[-1] if '|' in header.split()[0] else header.split()[0]
                            ref_sequences[protein_name][species] = protein_id
                        else:
                            # Try NCBI format with species in brackets: [Genus species ...]
                            match = re.search(r'\[([^\]]+)\]', header)
                            if match:
                                species_full = match.group(1).strip()
                                # Extract only genus and species (first two words)
                                species_match = re.match(r'(\S+\s+\S+)', species_full)
                                if species_match:
                                    species = species_match.group(1).strip()
                                else:
                                    species = species_full
                                protein_id = header.split()[0]
                                ref_sequences[protein_name][species] = protein_id
    
    return ref_sequences


def find_reference_kegg(table_df, ref_sequences, protein_name, protein_kegg_map):
    """
    Find KEGG annotation for reference sequence(s) by matching species names.
    Only returns KEGGs that match the expected KEGG from protein_kegg.tsv.
    Returns set of reference KEGGs.
    """
    ref_keggs = set()
    
    # Get the expected KEGG from protein_kegg.tsv
    expected_kegg = protein_kegg_map.get(protein_name, None)
    if expected_kegg:
        print(f"  Expected KEGG from protein_kegg.tsv: {expected_kegg}")
    
    if protein_name not in ref_sequences:
        # If no reference sequences, use expected KEGG if available
        if expected_kegg:
            ref_keggs.add(expected_kegg)
            print(f"  No reference sequences found, using expected KEGG: {expected_kegg}")
        return ref_keggs
    
    ref_species = ref_sequences[protein_name]
    
    # Look for matching species in the table
    for species_name in ref_species.keys():
        # Convert spaces to underscores for matching (table uses underscores)
        species_underscore = species_name.replace(' ', '_')
        
        # Try exact match with underscores
        matches = table_df[table_df['species'] == species_underscore]
        
        # If no exact match, try contains (case insensitive)
        if matches.empty:
            matches = table_df[table_df['species'].str.contains(species_name, case=False, na=False, regex=False)]
        
        if not matches.empty:
            for _, row in matches.iterrows():
                kegg = row.get('kegg', 'False')
                if kegg and kegg != 'False':
                    # Validate against expected KEGG from protein_kegg.tsv
                    if expected_kegg and kegg != expected_kegg:
                        print(f"  SKIPPING reference: {species_name} -> KEGG: {kegg} (doesn't match expected {expected_kegg})")
                        continue
                    ref_keggs.add(kegg)
                    print(f"  Found reference: {species_name} ({species_underscore}) -> KEGG: {kegg}")
    
    # If no valid KEGGs found from reference species, fall back to expected KEGG
    if not ref_keggs and expected_kegg:
        ref_keggs.add(expected_kegg)
        print(f"  No matching KEGGs from reference species, using expected KEGG: {expected_kegg}")
    
    return ref_keggs


def highlight_matching_kegg(input_table, output_excel, ref_seq_dirs, protein_kegg_file):
    """
    Read protein table, find reference KEGG, and create Excel with highlighted rows.
    Uses protein_kegg.tsv to validate that highlighted KEGGs are correct.
    """
    # Get protein name from filename
    protein_name = Path(input_table).stem.replace('_table', '')
    print(f"\nProcessing {protein_name}...")
    
    # Read table
    df = pd.read_csv(input_table, sep='\t')
    
    # Load protein to KEGG mapping from protein_kegg.tsv
    protein_kegg_map = load_protein_kegg_mapping(protein_kegg_file)
    
    # Parse reference sequences
    ref_sequences = parse_fasta_headers(ref_seq_dirs)
    
    # Find reference KEGG(s) - validated against protein_kegg.tsv
    ref_keggs = find_reference_kegg(df, ref_sequences, protein_name, protein_kegg_map)
    
    if not ref_keggs:
        print(f"  No reference KEGG found for {protein_name}")
        # Still create Excel without highlighting
        df.to_excel(output_excel, index=False, sheet_name=protein_name[:31])
        return
    
    print(f"  Reference KEGG(s): {ref_keggs}")
    
    # Create Excel workbook
    wb = Workbook()
    ws = wb.active
    ws.title = protein_name[:31]  # Excel sheet name limit
    
    # Write headers
    for col_idx, col_name in enumerate(df.columns, 1):
        ws.cell(row=1, column=col_idx, value=col_name)
    
    # Define blue fill for highlighting
    blue_fill = PatternFill(start_color="ADD8E6", end_color="ADD8E6", fill_type="solid")
    
    # Write data and highlight matching rows
    highlighted_count = 0
    for row_idx, (_, row) in enumerate(df.iterrows(), 2):
        kegg = row.get('kegg', 'False')
        should_highlight = kegg in ref_keggs and kegg != 'False'
        
        if should_highlight:
            highlighted_count += 1
        
        for col_idx, value in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if should_highlight:
                cell.fill = blue_fill
    
    # Save workbook
    wb.save(output_excel)
    print(f"  Highlighted {highlighted_count} rows with matching KEGG")
    print(f"  Saved to: {output_excel}")


def main():
    parser = argparse.ArgumentParser(
        description='Highlight protein table rows with matching reference KEGG annotations'
    )
    parser.add_argument('--table', required=True,
                       help='Input protein table TSV file')
    parser.add_argument('--output', required=True,
                       help='Output Excel file path')
    parser.add_argument('--ref-seqs', nargs='+', 
                       default=[
                           '../../input/ref_seqs/arc',
                           '../../input/ref_seqs/bact',
                           '../../input/ref_seqs/euk'
                       ],
                       help='Directories containing reference sequences (default: arc, bact, euk)')
    parser.add_argument('--protein-kegg', 
                       default='../../input/kegg/protein_kegg.tsv',
                       help='Path to protein_kegg.tsv file for KEGG validation (default: ../../input/kegg/protein_kegg.tsv)')
    
    args = parser.parse_args()
    
    highlight_matching_kegg(args.table, args.output, args.ref_seqs, args.protein_kegg)


if __name__ == '__main__':
    main()
